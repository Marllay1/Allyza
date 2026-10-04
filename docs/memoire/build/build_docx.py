# -*- coding: utf-8 -*-
import os, re, sys
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(__file__))
import content1 as c1, content2 as c2, content3 as c3

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "Allyza_Memoire_M1_Messagerie.docx")
VIOLET = RGBColor(0x3B, 0x1F, 0x52)
GREY = RGBColor(0x6B, 0x65, 0x75)
FONT = "Times New Roman"

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21), Cm(29.7)
sec.left_margin, sec.right_margin, sec.top_margin, sec.bottom_margin = Cm(2.8), Cm(2.2), Cm(2.5), Cm(2.3)
sec.different_first_page_header_footer = True

# ------------------------------------------------------------------ styles
def set_font(style, size, bold=None, italic=None, color=None, name=FONT):
    style.font.name = name
    style.element.rPr.rFonts.set(qn("w:eastAsia"), name)
    style.font.size = Pt(size)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
    if color is not None: style.font.color.rgb = color

st = doc.styles
set_font(st["Normal"], 11.5)
st["Normal"].paragraph_format.space_after = Pt(6)
st["Normal"].paragraph_format.line_spacing = 1.18
st["Normal"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name, size, before, after in (("Heading 1", 17, 0, 14), ("Heading 2", 13.5, 14, 6), ("Heading 3", 12, 10, 4)):
    s = st[name]; set_font(s, size, bold=True, color=VIOLET)
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.keep_with_next = True
    s.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
st["Heading 3"].font.italic = True
set_font(st["Caption"], 10, bold=False, italic=True, color=GREY)
st["Caption"].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
st["Caption"].paragraph_format.space_after = Pt(10)
from docx.enum.style import WD_STYLE_TYPE
ft = st.add_style("FrontTitle", WD_STYLE_TYPE.PARAGRAPH); ft.base_style = st["Normal"]
set_font(ft, 17, bold=True, color=VIOLET); ft.paragraph_format.space_after = Pt(14); ft.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
for lvl in (1, 2, 3):
    try:
        t = st[f"TOC {lvl}"]
    except KeyError:
        t = st.add_style(f"TOC {lvl}", WD_STYLE_TYPE.PARAGRAPH); t.base_style = st["Normal"]
    set_font(t, 11 if lvl == 1 else 10.5, bold=(lvl == 1)); t.paragraph_format.left_indent = Cm(0.6 * (lvl - 1)); t.paragraph_format.space_after = Pt(2)
    t.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
try:
    tof = st["Table of Figures"]
except KeyError:
    tof = st.add_style("Table of Figures", WD_STYLE_TYPE.PARAGRAPH); tof.base_style = st["Normal"]
set_font(tof, 10.5); tof.paragraph_format.space_after = Pt(2); tof.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
for ls in ("List Bullet", "List Number"):
    set_font(st[ls], 11.5); st[ls].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY; st[ls].paragraph_format.space_after = Pt(3)

# ------------------------------------------------------------------ helpers
def shade(cell, hexfill):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto"); sh.set(qn("w:fill"), hexfill); tcPr.append(sh)

TBLPR_ORDER = ["tblStyle", "tblpPr", "tblOverlap", "bidiVisual", "tblStyleRowBandSize", "tblStyleColBandSize", "tblW", "jc", "tblCellSpacing", "tblInd", "tblBorders", "shd", "tblLayout", "tblCellMar", "tblLook", "tblCaption", "tblDescription"]
def put(tblPr, el):
    name = el.tag.split("}")[1]
    for old in tblPr.findall(qn("w:" + name)): tblPr.remove(old)
    idx = TBLPR_ORDER.index(name)
    for child in list(tblPr):
        cn = child.tag.split("}")[1]
        if cn in TBLPR_ORDER and TBLPR_ORDER.index(cn) > idx:
            child.addprevious(el); return
    tblPr.append(el)

def cell_margins(tbl, top=50, bottom=50, left=90, right=90):
    tblPr = tbl._tbl.tblPr
    m = OxmlElement("w:tblCellMar")
    for k, v in (("top", top), ("left", left), ("bottom", bottom), ("right", right)):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:w"), str(v)); e.set(qn("w:type"), "dxa"); m.append(e)
    put(tblPr, m)

def borders(tbl, color="B9AEC6", sz=4):
    tblPr = tbl._tbl.tblPr
    b = OxmlElement("w:tblBorders")
    for k in ("top", "left", "bottom", "right", "insideH", "insideV"):
        e = OxmlElement(f"w:{k}"); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), str(sz)); e.set(qn("w:space"), "0"); e.set(qn("w:color"), color); b.append(e)
    put(tblPr, b)

TOK = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
def add_runs(par, text, size=None, color=None, bold=None):
    for part in TOK.split(text):
        if not part: continue
        if part.startswith("**") and part.endswith("**"): r = par.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"): r = par.add_run(part[1:-1]); r.font.name = "Consolas"; r.font.size = Pt((size or 11.5) - 1)
        elif part.startswith("*") and part.endswith("*") and len(part) > 2: r = par.add_run(part[1:-1]); r.italic = True
        else: r = par.add_run(part)
        if size and not part.startswith("`"): r.font.size = Pt(size)
        if color is not None: r.font.color.rgb = color
        if bold: r.bold = True
    return par

def field(par, instr, cached=""):
    def mk(t):
        r = par.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f); return r
    mk("begin")
    r = par.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it)
    mk("separate")
    par.add_run(cached)
    mk("end")

def para(text="", style=None, align=None, size=None, color=None, bold=None, space_after=None, keep=False):
    p = doc.add_paragraph(style=style)
    if align is not None: p.alignment = align
    add_runs(p, text, size=size, color=color, bold=bold)
    if space_after is not None: p.paragraph_format.space_after = Pt(space_after)
    if keep: p.paragraph_format.keep_with_next = True
    return p

def page_break():
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

fig_n = tab_n = 0
captions_fig, captions_tab = [], []

def caption(kind, text):
    global fig_n, tab_n
    if kind == "Figure": fig_n += 1; n = fig_n; captions_fig.append(f"Figure {n} — {text}")
    else: tab_n += 1; n = tab_n; captions_tab.append(f"Tableau {n} — {text}")
    p = doc.add_paragraph(style="Caption")
    p.add_run(f"{kind} ")
    field(p, f"SEQ {kind} \\* ARABIC", str(n))
    p.add_run(f" — {text}")
    return p

def table(cap, header, rows, widths, font, label_col=False):
    if cap:
        cp = caption("Tableau", cap); cp.paragraph_format.keep_with_next = True; cp.paragraph_format.space_after = Pt(4)
    ncol = len(header) if header else len(rows[0])
    widths = widths or [16.0 / ncol] * ncol
    scale = 16.0 / sum(widths); widths = [w * scale for w in widths]
    t = doc.add_table(rows=0, cols=ncol); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t); cell_margins(t)
    def fill(row, vals, head=False):
        cells = row.cells
        for i, v in enumerate(vals):
            c = cells[i]; c.width = Cm(widths[i])
            p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.paragraph_format.line_spacing = 1.05; p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            add_runs(p, str(v), size=font, bold=head or (label_col and i == 0), color=RGBColor(255, 255, 255) if head else None)
            if head: shade(c, "3B1F52")
            elif label_col and i == 0: shade(c, "EFE8F4")
    if header:
        r = t.add_row(); fill(r, header, True)
        trPr = r._tr.get_or_add_trPr(); h = OxmlElement("w:tblHeader"); h.set(qn("w:val"), "true"); trPr.append(h)
    for ri, row in enumerate(rows):
        r = t.add_row(); fill(r, row)
        if ri == 0 and header:
            for c in t.rows[0].cells + r.cells:
                for pp in c.paragraphs: pp.paragraph_format.keep_with_next = True
        trPr = r._tr.get_or_add_trPr(); cs = OxmlElement("w:cantSplit"); cs.set(qn("w:val"), "true"); trPr.append(cs)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)
    return t

def note(label, text):
    t = doc.add_table(rows=1, cols=1); t.alignment = WD_TABLE_ALIGNMENT.CENTER; t.autofit = False
    borders(t, color="F4C4B2", sz=8); cell_margins(t, 90, 90, 140, 140)
    c = t.rows[0].cells[0]; c.width = Cm(16.0); shade(c, "FFF4EE")
    p = c.paragraphs[0]; p.paragraph_format.space_after = Pt(0); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    add_runs(p, f"**{label}.** {text}", size=10.5)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

def figure(path, cap, width):
    full = os.path.join(ROOT, path)
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.paragraph_format.keep_with_next = True; p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(full, width=Cm(width))
    caption("Figure", cap)

def bullets(items, style="List Bullet"):
    for it in items:
        add_runs(doc.add_paragraph(style=style), it)

H_ENTRIES = []  # (level, text)
def render(blocks):
    for b in blocks:
        k = b[0]
        if k == "h1":
            if b[2]: page_break()
            doc.add_heading(b[1], level=1)
        elif k == "h2": doc.add_heading(b[1], level=2)
        elif k == "h3": doc.add_heading(b[1], level=3)
        elif k == "p": para(b[1])
        elif k == "bullets": bullets(b[1])
        elif k == "numbers": bullets(b[1], "List Number")
        elif k == "note": note(b[1], b[2])
        elif k == "table": table(b[1], b[2], b[3], b[4], b[5])
        elif k == "fig": figure(b[1], b[2], b[3])
        elif k == "pagebreak": page_break()

def scan(blocks):
    for b in blocks:
        if b[0] == "h1": H_ENTRIES.append((1, b[1]))
        elif b[0] == "h2": H_ENTRIES.append((2, b[1]))

# ------------------------------------------------------------------ assemble blocks
def problems_blocks():
    out = []
    for title, status, fields in c2.PROBLEMS:
        out.append(H3(title)); out.append(P(f"*{status}*"))
        out.append(("table", None, None, [[a, b] for a, b in fields], [2.6, 13.4], 9))
    return out

def scen_blocks():
    out = []
    for title, fields in c3.SCEN:
        out.append(H3(title)); out.append(("table", None, None, [[a, b] for a, b in fields], [3.6, 12.4], 9))
    return out

from blocks import *  # noqa
ANNEX = [
    H1("ANNEXES"),
    H2("Annexe A — Matrice de vérification des fonctionnalités"),
    P("Classes : **A** réellement implémentée ; **B** partielle ou non validée de bout en bout ; **C** prévue non finalisée ; **D** proposition / perspective. « P » : fourni par la plateforme."),
    ("table", "Matrice de vérification : élément, classe, preuve dans le projet, technologie, limites", ["Élément", "Cl.", "Preuve dans le projet", "Technologie", "Limites"], c3.MATRIX, [4.2, 1.5, 4.0, 3.0, 3.8], 8),
    H2("Annexe B — Variables d’environnement (noms uniquement)"),
    P("Les valeurs ne figurent pas dans ce document. L’état de leur définition en production n’est pas vérifiable depuis le dépôt."),
    ("table", "Variables d’environnement utilisées par la messagerie", ["Variable", "Usage", "Portée"], c3.ENVVARS, [6.6, 6.4, 3.0], 8.5),
    H2("Annexe C — Événements d’instrumentation proposés"),
    P("**Proposition, non implémentée.**"),
    ("table", "Événements proposés pour l’observabilité des appels et des médias", ["Événement", "Contenu", "Métriques"], c3.EVENTS, [5.4, 8.2, 2.4], 8.5),
    H2("Annexe D — Commandes de reproduction des mesures"),
    B(["`npm run test:db` : rejoue les 10 migrations dans PGlite et exécute les contrôles d’isolation (96 contrôles).",
       "`npm run test:e2ee` : tests des primitives cryptographiques (22 contrôles).",
       "`node scripts/e2ee-bench.mjs` : banc de calcul ; résultats bruts dans docs/memoire/bench-e2ee.json.",
       "`npx tsc --noEmit`, `npx eslint src`, `npx next build` : vérifications statiques et compilation."]),
    H2("Annexe E — Chronologie du dépôt"),
    ("table", "Jalons du dépôt (34 commits entre le 21 et le 25 septembre 2026)", ["Date", "Jalon (messages de commit résumés)"], [
        ["21/09", "Application PWA bilingue ; espaces personnel, Refuge et couple ; interface en verre dépoli ; tests de bout en bout"],
        ["23/09", "Messagerie (cadre sous clavier, appui long, surnoms) ; appels audio et vidéo WebRTC ; remplacement du service d’appel payant par TURN Cloudflare / coturn ; Web Push ; verrou d’application ; correctifs vocaux, haut-parleur, latence ; photos enregistrables ; tonalités"],
        ["25/09", "Persistance et palette d’appel ; historique et rappel ; chiffrement de bout en bout ; émojis et réactions par sticker ; recherche et liens ; correctifs de mise en page"],
    ], [2.2, 13.8], 9),
]

BODY = (c1.INTRO + c1.CH1 + c1.CH2 + c2.CH3 + c2.CH4 + problems_blocks() +
        c3.CH5 + scen_blocks() + c3.CH5B + c3.CH6 + c3.CONCLUSION)
scan(BODY)
H_ENTRIES.append((1, "BIBLIOGRAPHIE")); scan(ANNEX)

# pre-scan captions for the lists (so cached results are filled)
def prescan(blocks):
    f, t = [], []
    for b in blocks:
        if b[0] == "fig": f.append(b[2])
        if b[0] == "table" and b[1]: t.append(b[1])
    return f, t
f1, t1 = prescan(BODY + ANNEX)
PRE_F = [f"Figure {i+1} — {x}" for i, x in enumerate(f1)]
PRE_T = [f"Tableau {i+1} — {x}" for i, x in enumerate(t1)]

# ------------------------------------------------------------------ cover
logo = os.path.join(ROOT, "..", "..", "public", "brand", "lockup-light-560.png")
para("[ÉTABLISSEMENT — à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=14, bold=True, color=VIOLET, space_after=2)
para("[Département / UFR — à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=11, color=GREY, space_after=26)
para("MASTER 1", align=WD_ALIGN_PARAGRAPH.CENTER, size=16, bold=True, space_after=2)
para("Systèmes, réseaux, télécommunications et cybersécurité", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=2)
para("[Intitulé exact de la formation — à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY, space_after=40)
para("MÉMOIRE", align=WD_ALIGN_PARAGRAPH.CENTER, size=13, bold=True, color=VIOLET, space_after=10)
para("THÈME", align=WD_ALIGN_PARAGRAPH.CENTER, size=10, color=GREY, space_after=4)
para(c1.THEME, align=WD_ALIGN_PARAGRAPH.CENTER, size=20, bold=True, color=VIOLET, space_after=10)
para(c1.SUBTITLE, align=WD_ALIGN_PARAGRAPH.CENTER, size=13, space_after=30)
if os.path.exists(logo):
    p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER; p.add_run().add_picture(logo, width=Cm(6.0)); p.paragraph_format.space_after = Pt(30)
para("**Présenté par :** [Nom et prénom — à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=4)
para("**Sous l’encadrement de :** [Encadreur(s) — à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=4)
para("**Année académique :** [à compléter]", align=WD_ALIGN_PARAGRAPH.CENTER, size=12, space_after=4)

# ------------------------------------------------------------------ front matter
page_break(); para("Remerciements", style="FrontTitle")
para("[Page à rédiger par l’auteur : remerciements à l’encadreur, à l’établissement, aux relecteurs et à l’entourage. Aucun nom n’est inséré ici afin de ne rien présumer.]")
page_break(); para("Résumé", style="FrontTitle")
for t_ in c1.RESUME: para(t_)
para(f"**Mots-clés :** {c1.KEYWORDS_FR}.")
page_break(); para("Abstract", style="FrontTitle")
for t_ in c1.ABSTRACT: para(t_)
para(f"**Keywords:** {c1.KEYWORDS_EN}.")
page_break(); para("Liste des acronymes", style="FrontTitle")
table(None, ["Sigle", "Signification"], [[a, b] for a, b in c1.ACRONYMS], [3.0, 13.0], 10)

def toc_list(title, instr, entries, style, indent_fn=None):
    page_break(); para(title, style="FrontTitle")
    first = True
    for i, (lvl, text) in enumerate(entries):
        p = doc.add_paragraph(style=style(lvl) if callable(style) else style)
        if first:
            def mk(t):
                r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), t); r._r.append(f)
            mk("begin"); r = p.add_run(); it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r._r.append(it); mk("separate")
            first = False
        p.add_run(text)
        if i == len(entries) - 1:
            r = p.add_run(); f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), "end"); r._r.append(f)

toc_list("Table des matières", 'TOC \\o "1-2" \\h \\z \\u', H_ENTRIES, lambda l: f"TOC {l}")
toc_list("Liste des figures", 'TOC \\h \\z \\c "Figure"', [(1, x) for x in PRE_F], "Table of Figures")
toc_list("Liste des tableaux", 'TOC \\h \\z \\c "Tableau"', [(1, x) for x in PRE_T], "Table of Figures")
p = para("*Les numéros de page de ces listes sont calculés par Word : à l’ouverture, accepter la mise à jour des champs (ou Ctrl+A puis F9).*", size=9.5, color=GREY)

# ------------------------------------------------------------------ body
render(BODY)
page_break(); doc.add_heading("BIBLIOGRAPHIE", level=1)
for b in c3.BIB:
    p = para(b, size=10.5); p.alignment = WD_ALIGN_PARAGRAPH.LEFT; p.paragraph_format.left_indent = Cm(1.0); p.paragraph_format.first_line_indent = Cm(-1.0); p.paragraph_format.space_after = Pt(3)
note("Remarque de vérification", c3.BIB_NOTE.replace("Remarque de vérification : ", ""))
render(ANNEX)

# ------------------------------------------------------------------ header/footer
hdr = sec.header.paragraphs[0]; hdr.text = ""
r = hdr.add_run("Architecture de communication temps réel sécurisée et résiliente — Cas d’étude Allyza"); r.font.size = Pt(8.5); r.font.color.rgb = GREY; r.italic = True
hdr.alignment = WD_ALIGN_PARAGRAPH.LEFT
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
field(fp, "PAGE", "1")
for r in fp.runs: r.font.size = Pt(9.5)

# updateFields on open + core properties
settings = doc.settings.element
uf = OxmlElement("w:updateFields"); uf.set(qn("w:val"), "true")
anchor = next((settings.find(qn("w:" + n)) for n in ("hdrShapeDefaults", "footnotePr", "endnotePr", "compat", "docVars", "rsids") if settings.find(qn("w:" + n)) is not None), None)
if anchor is not None: anchor.addprevious(uf)
else: settings.append(uf)
z = settings.find(qn("w:zoom"))
if z is not None and z.get(qn("w:percent")) is None: z.set(qn("w:percent"), "100")
cp = doc.core_properties
cp.title = c1.THEME; cp.subject = c1.SUBTITLE; cp.author = ""; cp.keywords = c1.KEYWORDS_FR; cp.language = "fr-FR"
doc.save(OUT)
print("saved", OUT, "figures", fig_n, "tables", tab_n)
