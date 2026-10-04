# -*- coding: utf-8 -*-
import os, sys
from PIL import ImageFont
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Emu, Inches, Pt

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
OUT = os.path.join(ROOT, "Allyza_Soutenance_M1_Messagerie.pptx")
FIG = os.path.join(ROOT, "fig")
BRAND = os.path.abspath(os.path.join(ROOT, "..", "..", "public", "brand"))

def rgb(h): return RGBColor.from_string(h)
DEEP, VIOLET, TINT, PEACH, ROSE, TEAL, TEAL_T, INK, MUTED, WHITE, WARN = "24123F", "3B1F52", "F3EDF8", "F4C4B2", "C45C7E", "2F7F7A", "E3F1EF", "1E1A24", "6B6575", "FFFFFF", "FFF4EE"
HEAD, BODY = "Cambria", "Calibri"
W, H = 13.333, 7.5
MX = 0.6

prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
_fonts = {}
def font(bold, size):
    key = (bold, round(size * 4))
    if key not in _fonts:
        _fonts[key] = ImageFont.truetype("C:/Windows/Fonts/calibrib.ttf" if bold else "C:/Windows/Fonts/calibri.ttf", int(round(size * 4)))
    return _fonts[key]

def text_height(paragraphs, width_in, size, bold=False, factor=1.0, space=0.0):
    """Estimated height (inches) of paragraphs wrapped in width_in at `size` pt."""
    total = 0.0
    f = font(bold, size)
    for ptxt in paragraphs:
        words = ptxt.split(" "); lines = 1; cur = ""
        for w_ in words:
            trial = (cur + " " + w_).strip()
            if f.getlength(trial) / 4 * factor > width_in * 72 - 2:
                lines += 1; cur = w_
            else: cur = trial
        total += lines * size * 1.2 / 72 + space / 72
    return total

_slide_no = [0]
FITLOG = []
def new_slide(bg=WHITE, number=True):
    s = prs.slides.add_slide(BLANK); _slide_no[0] += 1
    s.background.fill.solid(); s.background.fill.fore_color.rgb = rgb(bg)
    s._number_on = number
    return s

def rect(s, x, y, w, h, fill=None, line=None, shape=MSO_SHAPE.RECTANGLE, lw=1.0, dash=False, radius=0.08, name=None):
    sh = s.shapes.add_shape(shape, Inches(x), Inches(y), Inches(w), Inches(h))
    if shape == MSO_SHAPE.ROUNDED_RECTANGLE: sh.adjustments[0] = radius
    if fill: sh.fill.solid(); sh.fill.fore_color.rgb = rgb(fill)
    else: sh.fill.background()
    if line:
        sh.line.color.rgb = rgb(line); sh.line.width = Pt(lw)
        if dash:
            from pptx.enum.dml import MSO_LINE
            sh.line.dash_style = MSO_LINE.DASH
    else: sh.line.fill.background()
    sh.shadow.inherit = False
    if name: sh.name = name
    return sh

def text(s, x, y, w, h, paras, size=16, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, face=BODY,
         bullet=False, space=4, min_size=12, name=None, check=True, italic=False):
    """paras: list of str or (str, dict(bold,color,size)). Auto-shrinks (to min_size) when the estimated height overflows."""
    plain = [p if isinstance(p, str) else p[0] for p in paras]
    sz = size
    inner_w = w - 0.2 - (0.25 if bullet else 0)
    while check and text_height(plain, inner_w, sz, bold, 1.04 if face == HEAD else 1.0, space) > h - 0.12 and sz > min_size:
        sz -= 1
    if check and text_height(plain, inner_w, sz, bold, 1.04 if face == HEAD else 1.0, space) > h - 0.12:
        FITLOG.append(f"slide {_slide_no[0]}: texte possiblement trop long ({plain[0][:40]!r}...)")
    tb = s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tb.text_frame.word_wrap = True
    tf = tb.text_frame; tf.margin_left = tf.margin_right = Inches(0.1); tf.margin_top = tf.margin_bottom = Inches(0.05); tf.vertical_anchor = anchor
    tb.name = name or "texte"
    for i, p in enumerate(paras):
        t, o = (p, {}) if isinstance(p, str) else p
        par = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        par.alignment = align; par.space_after = Pt(space)
        r = par.add_run(); r.text = (("• " if bullet else "") + t)
        r.font.name = face; r.font.size = Pt(o.get("size", sz)); r.font.bold = o.get("bold", bold); r.font.italic = o.get("italic", italic)
        r.font.color.rgb = rgb(o.get("color", color))
    return tb

def title(s, t, color=VIOLET, sub=None):
    text(s, MX, 0.35, W - 2 * MX, 0.9, [t], size=32, color=color, bold=True, face=HEAD, anchor=MSO_ANCHOR.MIDDLE, name="titre", min_size=26)
    if sub: text(s, MX, 1.2, W - 2 * MX, 0.45, [sub], size=16, color=MUTED, name="sous-titre", min_size=14)

def footer(s, dark=False):
    if not getattr(s, "_number_on", True): return
    c = "CFC4DB" if dark else MUTED
    text(s, W - 1.5, H - 0.45, 0.9, 0.3, [str(_slide_no[0])], size=12, color=c, align=PP_ALIGN.RIGHT, name="numéro", check=False)
    text(s, MX, H - 0.45, 8, 0.3, ["Allyza · Architecture de communication temps réel · Master 1"], size=12, color=c, name="pied", check=False)

def card(s, x, y, w, h, head, body, fill=TINT, head_color=VIOLET, body_size=17, head_size=20, tag=None, tag_color=TEAL):
    rect(s, x, y, w, h, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE, name=f"carte {head[:20]}")
    text(s, x + 0.15, y + 0.1, w - 0.3, 0.6, [head], size=head_size, bold=True, color=head_color, min_size=14, name="carte-titre")
    text(s, x + 0.15, y + 0.72, w - 0.3, h - 0.82, body if isinstance(body, list) else [body], size=body_size, min_size=12, space=3, name="carte-texte")
    if tag:
        rect(s, x + w - 0.75, y + h - 0.5, 0.55, 0.32, tag_color, shape=MSO_SHAPE.ROUNDED_RECTANGLE, radius=0.5)
        text(s, x + w - 0.8, y + h - 0.52, 0.65, 0.36, [tag], size=13, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, check=False)

def pic(s, path, x, y, w=None, h=None, name=None):
    p = s.shapes.add_picture(path, Inches(x), Inches(y), width=Inches(w) if w else None, height=Inches(h) if h else None)
    if name: p.name = name
    return p

def fit_pic(s, path, x, y, bw, bh):
    from PIL import Image
    iw, ih = Image.open(path).size; r = iw / ih
    w = bw; h = w / r
    if h > bh: h = bh; w = h * r
    return pic(s, path, x + (bw - w) / 2, y + (bh - h) / 2, w, h)

def placeholder(s, x, y, w, h, label):
    rect(s, x, y, w, h, "FBF9FD", line="B9AEC6", shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=True, radius=0.04, name="capture à insérer")
    text(s, x + 0.1, y + h / 2 - 0.4, w - 0.2, 0.8, ["Capture d’écran à insérer", (label, {"size": 12, "color": MUTED})], size=14, color=MUTED, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=12, check=False)

def notes(s, t): s.notes_slide.notes_text_frame.text = t

def arrow(s, x1, y1, x2, y2, color=VIOLET, w=1.75):
    c = s.shapes.add_connector(1, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
    c.line.color.rgb = rgb(color); c.line.width = Pt(w)
    ln = c.line._get_or_add_ln()
    from lxml import etree
    tail = etree.SubElement(ln, "{http://schemas.openxmlformats.org/drawingml/2006/main}tailEnd"); tail.set("type", "triangle")
    return c

def circle_num(s, x, y, n, d=0.5, fill=VIOLET):
    c = rect(s, x, y, d, d, fill, shape=MSO_SHAPE.OVAL)
    text(s, x - 0.05, y, d + 0.1, d, [str(n)], size=14, bold=True, color=WHITE, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, check=False)

# =========================================================================== 1. Titre
s = new_slide(DEEP, number=False)
text(s, MX, 0.7, 8.4, 0.4, ["MASTER 1 · SYSTÈMES, RÉSEAUX, TÉLÉCOMMUNICATIONS ET CYBERSÉCURITÉ"], size=14, color=PEACH, bold=True, check=False)
text(s, MX, 1.25, 8.8, 3.1, ["Conception d’une architecture de communication temps réel sécurisée et résiliente pour les applications multimédias distribuées"], size=32, color=WHITE, bold=True, face=HEAD, min_size=28, name="titre")
text(s, MX, 4.6, 8.6, 0.6, ["Cas d’étude : le module Messagerie du prototype expérimental Allyza"], size=20, color="E9E1F0", min_size=16)
text(s, MX, 5.5, 8.6, 1.3, ["Présenté par : [Nom et prénom — à compléter]", "Encadreur(s) : [à compléter]", "[Établissement — à compléter] · Année académique [à compléter]"], size=16, color="CFC4DB", space=3, min_size=14)
pic(s, os.path.join(BRAND, "lockup-dark-560.png"), 9.4, 1.7, h=3.1, name="logo Allyza")
notes(s, "Titre. Le sujet scientifique est l'architecture de communication temps réel ; Allyza est le prototype expérimental. Voir introduction générale du mémoire.")

# =========================================================================== 2. Contexte
s = new_slide(); title(s, "Des communications multimédias sur des réseaux peu fiables")
xs = [0.9, 5.0, 9.1]
labels = [("Utilisateurs", "Téléphones, navigateurs, PWA"), ("Réseaux hétérogènes", "Wi-Fi, 4G, 5G, NAT, coupures"), ("Services temps réel", "Signalisation, relais, stockage")]
for i, (h_, b_) in enumerate(labels):
    card(s, xs[i], 1.6, 3.3, 1.7, h_, b_, fill=TINT if i != 1 else WARN, body_size=17)
    if i < 2: arrow(s, xs[i] + 3.3, 2.45, xs[i + 1], 2.45)
text(s, MX, 3.65, 6.0, 3.3, [
    ("Pourquoi c’est difficile", {"bold": True, "color": VIOLET, "size": 22}),
    "La voix et la vidéo tolèrent peu de retard : environ 150 ms unidirectionnel (ITU-T G.114).",
    "La connectivité dépend de NAT, de changements de réseau et de coupures non annoncées.",
    "Le contenu est privé : sa compromission est irréversible."], size=18, space=9)
text(s, 6.9, 3.65, 5.8, 3.3, [
    ("Ce que le mémoire étudie", {"bold": True, "color": VIOLET, "size": 22}),
    "Les compromis entre sécurité, performance et disponibilité.",
    "Une architecture réellement implémentée, distinguée d’une architecture cible.",
    "Ce qui est mesuré, et ce qui reste à mesurer."], size=18, space=9)
footer(s); notes(s, "Contexte et justification — introduction générale, 1; chapitre 1, sections 1.1 à 1.4.")

# =========================================================================== 3. Problématique
s = new_slide(DEEP)
text(s, MX, 0.5, 6, 0.4, ["PROBLÉMATIQUE"], size=14, color=PEACH, bold=True, check=False)
text(s, MX, 1.1, W - 2 * MX, 2.5, ["Comment concevoir une architecture permettant de garantir simultanément confidentialité, intégrité, disponibilité et faible latence dans une plateforme de communication multimédia temps réel fonctionnant sur des réseaux hétérogènes ?"], size=30, color=WHITE, bold=True, face=HEAD, min_size=26, name="problématique")
chips = [("Sécurité", "confidentialité, intégrité"), ("Latence", "établissement, transmission"), ("Disponibilité", "continuité d’appel"), ("Résilience", "coupures, changements")]
for i, (a, b) in enumerate(chips):
    x = MX + i * 3.05
    rect(s, x, 4.5, 2.85, 1.7, "3B1F52", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.1, 4.62, 2.65, 0.6, [a], size=26, bold=True, color=PEACH, face=HEAD, min_size=20)
    text(s, x + 0.1, 5.3, 2.65, 0.8, [b], size=18, color="E9E1F0", min_size=14)
footer(s, dark=True); notes(s, "Problématique centrale. Elle reste le fil conducteur ; la conclusion y répond sans dépasser les preuves.")

# =========================================================================== 4. Objectifs
s = new_slide(); title(s, "Objectif général et six objectifs spécifiques")
rect(s, MX, 1.5, W - 2 * MX, 0.95, VIOLET, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, MX + 0.2, 1.5, W - 2 * MX - 0.4, 0.95, ["Concevoir et étudier une architecture de communication temps réel sécurisée et résiliente, à travers le cas d’étude Allyza."], size=18, color=WHITE, anchor=MSO_ANCHOR.MIDDLE, min_size=15)
objs = [("Concevoir", "reconstituer l’architecture réelle, puis proposer une cible"), ("Sécuriser", "authentification, RLS, stockage privé, E2EE optionnel"),
        ("Échanger en temps réel", "signalisation, WebRTC, médias"), ("Gérer les coupures", "reprise, rattrapage, persistance d’appel"),
        ("Évaluer", "tests automatisés, banc de calcul, protocole"), ("Étudier la résilience", "mécanismes et limites de la PWA")]
for i, (a, b) in enumerate(objs):
    x = MX + (i % 3) * 4.07; y = 2.8 + (i // 3) * 1.95
    card(s, x, y, 3.9, 1.75, a, b, body_size=17)
footer(s); notes(s, "Objectifs — introduction générale, 4.")

# =========================================================================== 5. Allyza
s = new_slide(); title(s, "Allyza : un prototype de messagerie privée (PWA)")
placeholder(s, 8.3, 1.55, 4.4, 5.2, "écran Messagerie (conversation)")
text(s, MX, 1.55, 7.4, 5.2, [
    ("Fait observé dans le dépôt", {"bold": True, "color": VIOLET, "size": 18}),
    "PWA bilingue pour deux personnes : une seule conversation par couple.",
    "Next.js 16, React 19, TypeScript ; Supabase (PostgreSQL, Auth, Storage, Realtime) ; Vercel.",
    "34 commits entre le 21 et le 25 septembre 2026.",
    ("Pourquoi un bon cas d’étude", {"bold": True, "color": VIOLET, "size": 18}),
    "Trois familles de trafic dans un même module : transactionnel, médias différés, temps réel.",
    "Contraintes réelles d’une PWA mobile : arrière-plan, notifications, audio verrouillé."], size=18, space=9)
footer(s); notes(s, "Présentation d'Allyza — chapitre 2, 2.1 et 2.2. Les captures d'écran sont à insérer par l'auteur : aucune capture exportée n'était disponible dans le dépôt.")

# =========================================================================== 6. Périmètre
s = new_slide(); title(s, "Périmètre : le module Messagerie uniquement")
rect(s, 4.9, 2.9, 3.5, 1.6, VIOLET, shape=MSO_SHAPE.OVAL)
text(s, 4.9, 2.9, 3.5, 1.6, ["Module", "Messagerie"], size=22, bold=True, color=WHITE, face=HEAD, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, check=False, space=0)
sats = [("Messages", "texte, stickers, réactions", 0.7, 1.55), ("Médias", "photos, documents (D)", 0.7, 3.35), ("Messages vocaux", "enregistrement, lecture", 0.7, 5.15),
        ("Appels audio", "WebRTC, sonnerie", 4.85, 1.45), ("Appels vidéo", "caméra, palette réduite", 8.9, 1.55), ("Reconnexion", "ICE restart, rattrapage", 8.9, 3.35),
        ("Sécurité", "RLS, E2EE (B), verrou", 8.9, 5.15), ("Notifications", "Web Push, pastilles", 4.85, 5.3)]
for a, b, x, y in sats:
    card(s, x, y, 3.75, 1.35, a, b, body_size=16, head_size=18)
footer(s); notes(s, "Délimitation — introduction générale, 6. Les fonctionnalités de classe D (documents, livraison, CallKit) ne sont pas présentées comme implémentées.")

# =========================================================================== 7. Besoins
s = new_slide(); title(s, "Besoins fonctionnels et non fonctionnels")
card(s, MX, 1.5, 5.95, 5.3, "Besoins fonctionnels", ["Messages texte, réponses, édition, suppression", "Stickers, émojis, réactions", "Photos et messages vocaux", "Appels audio et vidéo, sonnerie, historique", "Poursuivre l’appel hors de l’écran d’appel", "Notifications hors application", "Persistance et rattrapage"], body_size=19)
card(s, 6.78, 1.5, 5.95, 5.3, "Besoins non fonctionnels", ["Confidentialité et intégrité", "Disponibilité de l’appel et résilience", "Faible latence d’établissement", "Persistance des échanges", "Ergonomie mobile", "Notifications sans contenu"], body_size=19, fill=TEAL_T, head_color=TEAL)
footer(s); notes(s, "Besoins — chapitre 2, 2.5 et 2.6 (références BF1 à BF12, BN1 à BN8).")

# =========================================================================== 8. Technologies
s = new_slide(); title(s, "Technologies étudiées et choix retenus")
rows = [("WebSocket / Realtime", "Événements poussés", "postgres_changes + diffusion privée"), ("WebRTC", "Média pair-à-pair chiffré (DTLS-SRTP)", "RTCPeerConnection, offre/réponse"),
        ("Signalisation", "Hors norme : à fournir", "Canal de diffusion + table calls"), ("STUN / TURN / ICE", "Traversée de NAT, relais", "TURN Cloudflare (identifiants éphémères)"),
        ("HTTPS / TLS", "Transit seulement", "Assuré par Vercel et Supabase"), ("Web Push", "Notifier hors application", "VAPID, Service Worker, after()"),
        ("Services cloud", "Backend géré", "Supabase + Vercel + TURN")]
tbl = s.shapes.add_table(len(rows) + 1, 3, Inches(MX), Inches(1.5), Inches(W - 2 * MX), Inches(0.62 * (len(rows) + 1))).table
for j, hd in enumerate(["Technologie / concept", "Rôle", "Choix dans Allyza"]):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(VIOLET); c.text_frame.text = hd
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(15); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = rgb(WHITE)
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(TINT if i % 2 else WHITE); c.text_frame.text = v
        r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.name = BODY; r.font.color.rgb = rgb(INK); r.font.bold = (j == 0)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
for j, wd in enumerate([3.6, 4.3, 4.23]): tbl.columns[j].width = Inches(wd)
footer(s); notes(s, "État de l'art et choix — chapitre 1, 1.5 ; chapitre 3, 3.15.")

# =========================================================================== 9. Architecture générale
s = new_slide(); title(s, "Architecture actuelle : composants présents dans le dépôt")
fit_pic(s, os.path.join(FIG, "architecture-actuelle.png"), 0.4, 1.4, 8.6, 5.5)
text(s, 9.1, 1.5, 3.7, 5.2, [
    ("Quatre types de flux", {"bold": True, "color": VIOLET, "size": 18}),
    "Données : HTTPS vers l’application serveur, SQL sous RLS.",
    "Signalisation : diffusion privée Supabase.",
    "Média : pair-à-pair, relais TURN si besoin.",
    "Notifications : Web Push.",
    ("Pas de serveur de signalisation ni de serveur média propre.", {"italic": True, "color": MUTED})], size=17, space=9)
footer(s); notes(s, "Architecture générale — chapitre 3, 3.2. Cloudflare n'intervient que pour le relais TURN (3.5.1) : ce n'est pas une architecture cloud-native complète.")

# =========================================================================== 10. Messagerie
s = new_slide(); title(s, "Module Messagerie : trois chemins pour trois trafics")
cols = [("Transactionnel", "Texte, réactions, lecture", ["Action serveur (zod, débit)", "INSERT sous RLS", "postgres_changes", "Rattrapage à la reprise"], TINT, VIOLET),
        ("Médias différés", "Photos, vocaux", ["Redimensionnement, EXIF retiré", "Téléversement direct", "Chemin validé côté serveur", "URL signées 1 h"], TINT, VIOLET),
        ("Temps réel", "Appels audio et vidéo", ["Offre/réponse par diffusion", "ICE + TURN", "SRTP pair-à-pair", "ICE restart ≤ 3"], TEAL_T, TEAL)]
for i, (a, b, items, fill, col) in enumerate(cols):
    x = MX + i * 4.07
    rect(s, x, 1.5, 3.9, 5.3, fill, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.2, 1.65, 3.5, 0.55, [a], size=26, bold=True, color=col, face=HEAD, min_size=20)
    text(s, x + 0.2, 2.25, 3.5, 0.45, [b], size=17, color=MUTED, min_size=14)
    text(s, x + 0.2, 3.0, 3.5, 3.7, items, size=19, space=12, bullet=True)
footer(s); notes(s, "Chapitre 3, 3.3 à 3.6. Les diagrammes de séquence figurent dans le mémoire (envoi, réception, média, appel, appel entrant).")

# =========================================================================== 11. Appel
s = new_slide(); title(s, "Établissement d’un appel : de l’offre au média")
steps = [("1", "Offre (SDP)", "Diffusée sur le canal privé ; renvoyée toutes les 3 s"), ("2", "Sonnerie", "L’offre porte le type d’appel : B sonne sans attendre la base"),
         ("3", "Réponse (SDP)", "Après acceptation et capture du micro"), ("4", "Candidats ICE", "Direct (host, srflx) ou relayé par TURN"), ("5", "Média", "DTLS puis SRTP, pair-à-pair")]
for i, (n, a, b) in enumerate(steps):
    x = MX + i * 2.45
    rect(s, x, 1.6, 2.25, 3.0, TEAL_T if i == 4 else TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    circle_num(s, x + 0.15, 1.75, n, 0.5, TEAL if i == 4 else VIOLET)
    text(s, x + 0.1, 2.35, 2.05, 0.5, [a], size=20, bold=True, color=VIOLET, min_size=16)
    text(s, x + 0.1, 2.9, 2.05, 1.65, [b], size=15, min_size=13)
    if i < 4: arrow(s, x + 2.25, 3.1, x + 2.45, 3.1)
text(s, MX, 4.95, 6.0, 1.9, [("Ce que fait le code", {"bold": True, "color": VIOLET, "size": 20}), "Un seul appel actif à la fois.", "Serveurs ICE en cache 2 h, préchargés.", "Table calls : garantie si la diffusion est perdue."], size=17, space=6)
text(s, 7.0, 4.95, 5.7, 1.9, [("Ce qui n’est pas mesuré", {"bold": True, "color": ROSE, "size": 20}), "Le temps d’établissement avant/après optimisation.", "Le taux de réussite entre réseaux différents."], size=17, space=6)
footer(s); notes(s, "Chapitre 3, 3.4 et 3.5 ; diagramme de séquence détaillé dans le mémoire (figure de la section 3.5) ; problème P5 (latence : optimisations livrées, effet non mesuré).")

# =========================================================================== 12. Sécurité
s = new_slide(); title(s, "Sécurité : des couches, et ce qui est réellement prouvé")
chain = [("Authentification", "Supabase Auth, JWT, limitation d’essais", "A"), ("Contrôle d’accès", "RLS : 70 politiques SQL du projet", "A"), ("Transit", "TLS/WSS (hébergeurs), DTLS-SRTP", "A"),
         ("Stockage privé", "Bucket privé, URL signées 1 h", "A"), ("E2EE optionnel", "ECDH + AES-GCM, MAC des empreintes", "B")]
for i, (a, b, c) in enumerate(chain):
    x = MX + i * 2.45
    card(s, x, 1.6, 2.25, 2.7, a, b, fill=TINT if c == "A" else WARN, body_size=15, head_size=17, tag=c, tag_color=TEAL if c == "A" else ROSE)
    if i < 4: arrow(s, x + 2.25, 2.95, x + 2.45, 2.95)
rect(s, MX, 4.6, W - 2 * MX, 2.2, WARN, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, MX + 0.2, 4.7, W - 2 * MX - 0.4, 2.0, [
    ("TLS ≠ chiffrement de bout en bout", {"bold": True, "color": VIOLET, "size": 20}),
    "Sans E2EE actif, l’opérateur de la base peut lire le contenu. L’E2EE est implémenté et testé unitairement (classe B) mais non validé entre deux comptes réels, et sa migration n’est pas confirmée en production.",
    "Pas de secret de transmission aval ; métadonnées visibles."], size=16, space=6)
footer(s); notes(s, "Chapitre 3, 3.10 ; chapitre 4, 4.12. A = implémenté et vérifié dans le code ; B = dans le code, non validé de bout en bout.")

# =========================================================================== 13. Résilience
s = new_slide(); title(s, "Résilience : des mécanismes, une efficacité à prouver")
phases = ["Communication normale", "Coupure ou changement de réseau", "Détection (connectionState)", "Reprise ICE (≤ 3 essais)", "Reprise ou abandon (20 s)"]
for i, t_ in enumerate(phases):
    x = MX + i * 2.45
    rect(s, x, 1.55, 2.25, 1.3, PEACH if i in (0, 4) else TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.05, 1.55, 2.15, 1.3, [t_], size=16, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=13)
    if i < 4: arrow(s, x + 2.25, 2.2, x + 2.45, 2.2)
for i, (a, b) in enumerate([("Signalisation double", "diffusion rapide + table durable"), ("Rattrapage", "relecture des messages à la reprise"),
                            ("Redémarrage ICE", "par l’appelant, abandon après 20 s"), ("Reprise du micro", "nouvelle capture, même connexion")]):
    card(s, MX + i * 3.05, 3.2, 2.85, 2.1, a, b, body_size=16, head_size=18)
text(s, MX, 5.65, W - 2 * MX, 1.2, [("Aucun temps de reprise n’a été mesuré : les scénarios S-5 et S-6 sont protocolés, pas exécutés.", {"bold": True, "color": ROSE})], size=20)
footer(s); notes(s, "Chapitre 3, 3.9 et 3.11 ; chapitre 5, 5.6. Implémenté (classe A pour le code) ≠ démontré par l'expérience.")

# =========================================================================== 14. Persistance d'appel
s = new_slide(); title(s, "Un appel est un état, pas un écran")
phs = ["idle", "sortant / entrant", "connecting", "connected", "ended"]
for i, t_ in enumerate(phs):
    x = MX + i * 1.7
    rect(s, x, 1.6, 1.5, 0.8, PEACH if t_ == "connected" else TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x, 1.6, 1.5, 0.8, [t_], size=14, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=12)
    if i < 4: arrow(s, x + 1.5, 2.0, x + 1.7, 2.0)
rect(s, MX, 2.9, 3.9, 2.3, VIOLET, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, MX + 0.15, 3.0, 3.6, 2.1, [("CallProvider", {"bold": True, "size": 20, "color": PEACH}), "Monté dans le layout. Porte la connexion, les flux, la signalisation, les minuteries."], size=16, color=WHITE, space=5)
rect(s, MX + 4.5, 2.9, 3.9, 2.3, TEAL_T, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
text(s, MX + 4.65, 3.0, 3.6, 2.1, [("CallOverlay", {"bold": True, "size": 20, "color": TEAL}), "Simple vue : plein écran ou palette. Peut disparaître sans toucher à l’appel."], size=16, space=5)
arrow(s, MX + 3.9, 4.05, MX + 4.5, 4.05, TEAL)
text(s, MX, 5.5, 8.4, 1.3, ["Retiré : le raccrochage sur pagehide (iOS le déclenche en arrière-plan). Ajouté : Wake Lock d’écran, reprise du micro."], size=16)
text(s, 9.4, 1.55, 3.4, 5.3, [
    ("Hypothèse H1", {"bold": True, "color": VIOLET, "size": 20}),
    "Séparer l’état de l’appel de son affichage le rend tolérant aux changements d’interface.",
    ("Limite de la PWA", {"bold": True, "color": ROSE, "size": 20}),
    "Le système peut suspendre la page : aucune API web ne le garantit."], size=16, space=8)
footer(s); notes(s, "Chapitre 3, 3.8 ; problème P4. Le comportement hors écran est traité comme un problème de gestion d'état, pas d'interface.")

# =========================================================================== 15. Implémentation
s = new_slide(); title(s, "Implémentation du prototype")
caps = ["Conversation", "Message vocal", "Appel audio", "Appel vidéo", "Palette d’appel"]
for i, c_ in enumerate(caps):
    x = MX + i * 2.45
    placeholder(s, x, 1.5, 2.3, 4.4, c_)
    text(s, x, 5.95, 2.3, 0.4, [c_], size=14, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, check=False)
text(s, MX, 6.35, W - 2 * MX, 0.6, ["Captures à insérer par l’auteur : aucune image exportée n’était disponible dans le dépôt au moment de la rédaction."], size=12, color=MUTED, check=False)
footer(s); notes(s, "Chapitre 4. Insérer des captures réelles de l'application (écran Messagerie, vocal, appels audio/vidéo, palette).")

# =========================================================================== 16. Problèmes
s = new_slide(); title(s, "Problèmes rencontrés, corrections et statut de preuve")
rows = [("Vocaux rejetés", "Bucket limité aux images", "Types audio autorisés (migration 0009)", "Vérifié par essai"),
        ("Appel audio muet", "Flux sans élément média", "Élément média invisible", "À valider sur appareils"),
        ("Coupure en arrière-plan", "Raccrochage sur pagehide", "Gestionnaire retiré, Wake Lock", "Limite iOS persistante"),
        ("Latence d’établissement", "ICE à la demande, attente base", "Cache + offre diffusée", "Non mesuré"),
        ("Push muet en production", "Promesse non attendue (serverless)", "Envoi dans after()", "À valider en production"),
        ("Appels qui ne s’établissent plus", "Non établie (TURN ? session audio ?)", "Diagnostics ajoutés", "Non confirmé résolu")]
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(MX), Inches(1.5), Inches(W - 2 * MX), Inches(0.7 * (len(rows) + 1))).table
for j, hd in enumerate(["Symptôme", "Cause", "Correction", "Statut"]):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(VIOLET); c.text_frame.text = hd
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(15); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = rgb(WHITE)
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.fill.solid()
        bad = (j == 3 and ("Non" in v or "À valider" in v or "Limite" in v))
        c.fill.fore_color.rgb = rgb(WARN if bad else (TINT if i % 2 else WHITE)); c.text_frame.text = v
        r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.name = BODY; r.font.color.rgb = rgb(INK); r.font.bold = (j == 0)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
for j, wd in enumerate([2.9, 3.4, 3.4, 2.43]): tbl.columns[j].width = Inches(wd)
footer(s); notes(s, "Chapitre 4, 4.14 (P1 à P13, format symptôme/cause/diagnostic/correction/impact/limites/validation). Aucun problème n'est déclaré résolu sans preuve.")

# =========================================================================== 17. Méthodologie / scénarios
s = new_slide(); title(s, "Méthodologie expérimentale : huit scénarios")
rows = [("S-1", "Réseau stable", "Établissement, taux de réussite, type de chemin", "Protocolé"), ("S-2", "Latence élevée", "Effet du délai sur l’établissement et la qualité", "Protocolé"),
        ("S-3", "Perte de paquets", "Tolérance de la voix et de la vidéo", "Protocolé"), ("S-4", "Bande passante limitée", "Adaptation, temps de chargement des médias", "Protocolé"),
        ("S-5", "Coupure temporaire", "Reprise de l’appel, rattrapage des messages", "Protocolé"), ("S-6", "Changement de réseau", "ICE restart Wi-Fi ↔ 4G", "Protocolé"),
        ("S-7", "Charge croissante", "Non pertinent pour deux utilisateurs", "Non réalisable"), ("S-8", "Sécurité", "Isolation SQL, cryptographie", "MESURÉ")]
tbl = s.shapes.add_table(len(rows) + 1, 4, Inches(MX), Inches(1.5), Inches(W - 2 * MX), Inches(0.58 * (len(rows) + 1))).table
for j, hd in enumerate(["Réf.", "Scénario", "Objectif", "État"]):
    c = tbl.cell(0, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(VIOLET); c.text_frame.text = hd
    r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(15); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = rgb(WHITE)
for i, row in enumerate(rows, 1):
    for j, v in enumerate(row):
        c = tbl.cell(i, j); c.fill.solid(); c.fill.fore_color.rgb = rgb(TEAL_T if row[3] == "MESURÉ" else (TINT if i % 2 else WHITE)); c.text_frame.text = v
        r = c.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.name = BODY; r.font.color.rgb = rgb(INK); r.font.bold = (j in (0, 3) and row[3] == "MESURÉ") or j == 0
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
for j, wd in enumerate([1.0, 3.3, 5.8, 2.03]): tbl.columns[j].width = Inches(wd)
footer(s); notes(s, "Chapitre 5, 5.1 à 5.4 : méthode But–Question–Métrique, au moins 30 répétitions par condition, médiane et p95. La limitation réseau des outils de développement du navigateur n'affecte pas le trafic WebRTC en UDP.")

# =========================================================================== 18. Résultats
s = new_slide(); title(s, "Résultats : ce qui a été réellement mesuré")
stats = [("96", "contrôles d’isolation SQL réussis, 0 échec"), ("22", "contrôles cryptographiques réussis"), ("< 1,2 ms", "chiffrer 4 000 caractères (médiane)"), ("+28 octets", "surcoût d’un fichier chiffré")]
for i, (a, b) in enumerate(stats):
    x = MX + i * 3.05
    rect(s, x, 1.45, 2.85, 1.65, TINT, shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.1, 1.5, 2.65, 0.75, [a], size=34, bold=True, color=VIOLET, face=HEAD, min_size=26, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE)
    text(s, x + 0.1, 2.25, 2.65, 0.8, [b], size=15, color=MUTED, align=PP_ALIGN.CENTER, min_size=12)
cd = CategoryChartData(); cd.categories = ["20 car.", "200 car.", "1 000 car.", "4 000 car."]
cd.add_series("Chiffrement", (0.23, 0.26, 0.50, 1.13)); cd.add_series("Déchiffrement", (0.20, 0.21, 0.27, 0.34))
gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MX), Inches(3.3), Inches(6.0), Inches(3.5), cd); ch = gf.chart
cd2 = CategoryChartData(); cd2.categories = ["50 ko", "300 ko", "3 Mo"]
cd2.add_series("Chiffrement", (0.61, 2.08, 19.1)); cd2.add_series("Déchiffrement", (0.57, 1.70, 18.5))
gf2 = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.75), Inches(3.3), Inches(6.0), Inches(3.5), cd2); ch2 = gf2.chart
for chart, ttl in ((ch, "Textes : temps de calcul (ms, médiane)"), (ch2, "Fichiers : temps de calcul (ms, médiane)")):
    chart.has_title = True; chart.chart_title.text_frame.text = ttl
    r = chart.chart_title.text_frame.paragraphs[0].runs[0]; r.font.size = Pt(14); r.font.bold = True; r.font.name = BODY; r.font.color.rgb = rgb(VIOLET)
    chart.has_legend = True; chart.legend.position = XL_LEGEND_POSITION.BOTTOM; chart.legend.include_in_layout = False; chart.legend.font.size = Pt(12); chart.legend.font.name = BODY
    chart.category_axis.tick_labels.font.size = Pt(12); chart.value_axis.tick_labels.font.size = Pt(12)
    chart.value_axis.has_major_gridlines = True; chart.value_axis.major_gridlines.format.line.color.rgb = rgb("E3DCEB")
    chart.category_axis.format.line.color.rgb = rgb("B9AEC6")
    for k, col in enumerate((VIOLET, PEACH)):
        ser = chart.plots[0].series[k]; ser.format.fill.solid(); ser.format.fill.fore_color.rgb = rgb(col)
    pl = chart.plots[0]; pl.has_data_labels = True; pl.data_labels.font.size = Pt(11); pl.data_labels.number_format = "0.0#"; pl.data_labels.number_format_is_linked = False
    pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END; pl.gap_width = 60
footer(s); notes(s, "Chapitre 5, 5.7. Mesures de calcul sur un poste de bureau (Node.js v26.1.0, Windows x64, WebCrypto) : ce ne sont pas des temps de bout en bout ni des mesures sur téléphone. Aucune mesure de latence, perte ou reprise réseau n'existe à ce jour.")

# =========================================================================== 19. Analyse
s = new_slide(); title(s, "Ce que les mesures apprennent, et ce qu’elles taisent")
card(s, MX, 1.5, 5.95, 5.3, "Enseignements", [
    "Le calcul de chiffrement est faible devant le budget de 150 ms (poste de test).",
    "Le vrai coût de l’E2EE est d’usage : mot de passe de chiffrement, déverrouillage d’environ 2 s, perte possible des données.",
    "L’isolation est exprimée dans la base, pas seulement dans l’interface."], body_size=18)
card(s, 6.78, 1.5, 5.95, 5.3, "Constat d’absence de preuve", [
    "Latence d’établissement, gigue, perte, temps de reprise, taux de réussite : non mesurés.",
    "Aucune instrumentation dans l’application actuelle.",
    "H2 partiellement testée ; H3 non testée.",
    "Mesures sur un poste de bureau, pas sur téléphone."], body_size=18, fill=WARN, head_color=ROSE)
footer(s); notes(s, "Chapitre 5, 5.8 : analyses 1 à 4 et constat d'absence de preuve.")

# =========================================================================== 20. Compromis
s = new_slide(); title(s, "Le triangle des compromis")
tri = s.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE, Inches(1.4), Inches(2.0), Inches(4.6), Inches(3.9))
tri.fill.solid(); tri.fill.fore_color.rgb = rgb(TINT); tri.line.color.rgb = rgb(VIOLET); tri.line.width = Pt(2.5); tri.shadow.inherit = False
text(s, 2.2, 1.5, 3.0, 0.5, ["SÉCURITÉ"], size=20, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, check=False)
text(s, 0.4, 5.95, 3.4, 0.8, ["PERFORMANCE", ("latence, qualité", {"size": 14, "bold": False, "color": MUTED})], size=18, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, space=0, check=False)
text(s, 3.9, 5.95, 3.4, 0.8, ["DISPONIBILITÉ", ("résilience", {"size": 14, "bold": False, "color": MUTED})], size=18, bold=True, color=VIOLET, align=PP_ALIGN.CENTER, space=0, check=False)
text(s, 7.5, 1.6, 5.3, 5.2, [
    ("Pas de « meilleure » solution", {"bold": True, "color": VIOLET, "size": 22}),
    "E2EE : protège de l’opérateur ; ajoute la gestion de clés et un risque de perte.",
    "Reprise automatique : améliore la continuité ; ajoute des états à gérer et à tester.",
    "Métadonnées minimales : protègent ; réduisent le diagnostic, donc l’observabilité.",
    "Relais TURN : fiabilise la connexion ; ajoute un saut réseau (non mesuré)."], size=18, space=10)
footer(s); notes(s, "Chapitre 6, 6.3 : tableau d'analyse des compromis avec l'état de la preuve de chaque mécanisme.")

# =========================================================================== 21. Limites
s = new_slide(); title(s, "Limites du prototype, en toute transparence")
lim = [("Contraintes PWA", ["Pas de CallKit/PushKit", "Arrière-plan décidé par le système", "Audio verrouillé avant le premier geste", "Pas de choix du haut-parleur sur Safari iOS"]),
       ("Temps réel", ["Relais TURN à configurer en production", "Appels un-à-un uniquement", "Aucune métrique de qualité collectée"]),
       ("Méthode", ["Aucune campagne sur appareils", "Tests SQL dans PGlite, pas sur Supabase réel", "E2EE non validé entre deux comptes", "Bibliographie à revérifier"])]
for i, (a, items) in enumerate(lim):
    card(s, MX + i * 4.07, 1.5, 3.9, 5.3, a, items, body_size=18, fill=TINT if i != 1 else WARN)
footer(s); notes(s, "Chapitre 6, 6.5.")

# =========================================================================== 22. Architecture cible
s = new_slide(); title(s, "Architecture cible : une proposition, non implémentée")
def dbox(x, y, w, h, t_, fill=TINT, col=VIOLET):
    rect(s, x, y, w, h, fill, line=col, shape=MSO_SHAPE.ROUNDED_RECTANGLE, dash=True, lw=1.5, radius=0.1)
    text(s, x, y, w, h, [t_], size=15, bold=True, color=INK, align=PP_ALIGN.CENTER, anchor=MSO_ANCHOR.MIDDLE, min_size=12)
dbox(MX, 1.5, 2.6, 0.95, "PWA + application iOS native", PEACH, ROSE); dbox(MX + 2.8, 1.5, 2.6, 0.95, "Passerelle API / Auth"); dbox(MX + 5.6, 1.5, 2.6, 0.95, "Observabilité", TEAL_T, TEAL)
dbox(MX, 3.0, 2.6, 0.95, "Service messagerie"); dbox(MX + 2.8, 3.0, 2.6, 0.95, "Service signalisation"); dbox(MX + 5.6, 3.0, 2.6, 0.95, "Service médias")
dbox(MX, 4.5, 2.6, 0.95, "Base répliquée, sauvegardes"); dbox(MX + 2.8, 4.5, 2.6, 0.95, "Stockage objet chiffré"); dbox(MX + 5.6, 4.5, 2.6, 0.95, "STUN / TURN redondants", TEAL_T, TEAL)
dbox(MX, 5.95, 8.2, 0.85, "Sécurité : E2EE à cliquet, rotation des clés, audit, tests d’intrusion", WARN, PEACH)
for x_ in (MX + 1.3, MX + 4.1, MX + 6.9):
    arrow(s, x_, 2.45, x_, 3.0); arrow(s, x_, 3.95, x_, 4.5)
text(s, 9.2, 1.5, 3.6, 5.3, [
    ("Évolutions proposées", {"bold": True, "color": VIOLET, "size": 20}),
    "Services séparés derrière une passerelle.", "Journaux, métriques, getStats().", "E2EE à cliquet, audité.", "Relais TURN multi-régions.", "iOS natif : CallKit, PushKit."], size=17, space=9)
footer(s); notes(s, "Chapitre 3, 3.12, et chapitre 6, 6.6. Tout ce qui figure ici est une perspective : ne pas présenter CallKit ou PushKit comme faisant partie de l'implémentation actuelle.")

# =========================================================================== 23. Perspectives
s = new_slide(); title(s, "Perspectives")
steps = [("Prototype actuel", "Architecture réelle documentée"), ("Mesurer", "Instrumenter, exécuter S-1 à S-6"), ("Renforcer", "E2EE à cliquet, relais redondants"), ("Étendre", "Natif iOS, groupes, charge")]
for i, (a, b) in enumerate(steps):
    x = MX + i * 3.05
    card(s, x, 2.0, 2.75, 2.8, a, b, body_size=18, fill=TINT if i else TEAL_T, head_color=VIOLET if i else TEAL)
    if i < 3: arrow(s, x + 2.75, 3.4, x + 3.05, 3.4)
text(s, MX, 5.3, W - 2 * MX, 1.4, ["Tout ce qui précède est une proposition ou une perspective, non confirmée comme implémentée."], size=20, color=ROSE, bold=True)
footer(s); notes(s, "Chapitre 6, 6.6 et 6.7.")

# =========================================================================== 24. Apports
s = new_slide(); title(s, "Apports du projet")
card(s, MX, 1.5, 5.95, 5.3, "Apports techniques", ["Séparation état d’appel / affichage", "Signalisation double : diffusion + base", "E2EE optionnel, refus de rétrogradation", "Problèmes réels et corrections documentés", "Protocole expérimental prêt à l’emploi"], body_size=19)
card(s, 6.78, 1.5, 5.95, 5.3, "Apports académiques", ["Réseaux : NAT, ICE, TURN, résilience", "Cybersécurité : RLS, TLS vs E2EE, métadonnées", "Systèmes distribués : cohérence à terme", "Cloud : services gérés et limites", "Temps réel : WebRTC, signalisation"], body_size=19, fill=TEAL_T, head_color=TEAL)
footer(s); notes(s, "Chapitre 6, 6.4.")

# =========================================================================== 25. Conclusion
s = new_slide(DEEP)
text(s, MX, 0.5, 6, 0.4, ["CONCLUSION"], size=14, color=PEACH, bold=True, check=False)
flow = [("Problème", "Garantir confidentialité, intégrité, disponibilité, faible latence"), ("Démarche", "Audit du dépôt, classes A/B/C/D, tests, protocole"), ("Solution", "Backend géré, WebRTC, E2EE optionnel"), ("Résultats", "96 + 22 contrôles ; chiffrement peu coûteux en calcul"), ("Perspectives", "Mesurer, renforcer, étendre")]
for i, (a, b) in enumerate(flow):
    x = MX + i * 2.45
    rect(s, x, 1.2, 2.3, 3.3, "3B1F52", shape=MSO_SHAPE.ROUNDED_RECTANGLE)
    text(s, x + 0.1, 1.3, 2.1, 0.55, [a], size=22, bold=True, color=PEACH, face=HEAD, min_size=17)
    text(s, x + 0.1, 1.95, 2.1, 2.4, [b], size=16, color="E9E1F0", min_size=13)
text(s, MX, 4.9, W - 2 * MX, 1.9, ["L’architecture proposée n’est pas démontrée comme garantissant simultanément les quatre propriétés ; elle est construisible, ses compromis sont explicites et ce qui reste à établir est mesurable par un protocole précis."], size=24, color=WHITE, face=HEAD, min_size=18)
footer(s, dark=True); notes(s, "Conclusion générale : répond à la problématique sans aller au-delà des preuves.")

# =========================================================================== 26. Merci
s = new_slide(DEEP, number=False)
pic(s, os.path.join(BRAND, "lockup-dark-560.png"), 5.2, 0.7, h=3.0, name="logo Allyza")
text(s, 0, 4.0, W, 0.9, ["Merci pour votre attention"], size=40, bold=True, color=WHITE, face=HEAD, align=PP_ALIGN.CENTER, check=False)
text(s, 0, 5.0, W, 0.6, ["Questions et échanges"], size=22, color=PEACH, align=PP_ALIGN.CENTER, check=False)
notes(s, "Questions.")

for sl in prs.slides: footer(sl, dark=False) if False else None
prs.save(OUT)
print("saved", OUT, "slides", len(prs.slides))
if FITLOG: print("FIT WARNINGS:"); [print(" -", f) for f in FITLOG]
